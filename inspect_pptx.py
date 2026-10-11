"""inspect_pptx.py: Kiểm tra cấu trúc, số lượng slide, tính chỉnh sửa (editability) và speaker notes của file PPTX."""

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    print(
        "Lỗi: Thư viện 'python-pptx' chưa được cài đặt. "
        "Vui lòng chạy: pip install python-pptx",
        file=sys.stderr,
    )
    sys.exit(1)


class PPTXInspector:
    """Bộ kiểm tra cấu trúc hình học và văn bản của tệp PPTX."""

    def __init__(self, pptx_path: str | Path):
        self.pptx_path = Path(pptx_path).resolve()
        if not self.pptx_path.exists():
            raise FileNotFoundError(f"Không tìm thấy tệp PPTX: {self.pptx_path}")

        try:
            self.prs = Presentation(str(self.pptx_path))
        except Exception as e:
            raise ValueError(f"Tệp không phải định dạng PowerPoint hợp lệ hoặc bị hỏng: {e}")

    def inspect_slide(self, slide_idx: int, slide) -> Dict[str, Any]:
        """Phân tích chi tiết từng slide: đếm shape, text box, table, ảnh và speaker notes."""
        slide_id = slide_idx + 1
        issues: List[str] = []

        text_frames_count = 0
        total_text_length = 0
        tables_count = 0
        table_cells_count = 0
        shapes_count = 0
        pictures_count = 0

        # 1. Duyệt qua toàn bộ shapes trên slide
        for shape in slide.shapes:
            shape_type = shape.shape_type

            # Kiểm tra bảng native (chỉnh sửa được)
            if shape.has_table:
                tables_count += 1
                table = shape.table
                table_cells_count += len(table.rows) * len(table.columns)
                for row in table.rows:
                    for cell in row.cells:
                        total_text_length += len(cell.text.strip())

            # Kiểm tra khung văn bản (native text box / placeholder)
            elif shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    text_frames_count += 1
                    total_text_length += len(text)

            # Đếm ảnh phẳng
            if shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures_count += 1
            elif shape_type in (
                MSO_SHAPE_TYPE.AUTO_SHAPE,
                MSO_SHAPE_TYPE.FREEFORM,
                MSO_SHAPE_TYPE.GROUP,
            ):
                shapes_count += 1

        # 2. Kiểm tra speaker notes
        has_speaker_notes = False
        speaker_notes_text = ""
        if slide.has_notes_slide:
            notes_slide = slide.notes_slide
            if notes_slide.notes_text_frame:
                speaker_notes_text = notes_slide.notes_text_frame.text.strip()
                has_speaker_notes = len(speaker_notes_text) > 0

        # 3. Chốt chặn kiểm định: Phát hiện rasterization (ảnh phẳng lấp liếm slide)
        # Nếu slide không có bất kỳ khối chữ hoặc bảng nào, nhưng lại có ảnh
        is_rasterized = False
        if text_frames_count == 0 and tables_count == 0 and pictures_count > 0:
            is_rasterized = True
            issues.append(
                "Slide bị nghi ngờ raster hóa: Toàn bộ nội dung là ảnh phẳng, không có text box hoặc bảng chỉnh sửa được."
            )

        # 4. Kiểm tra slide rỗng
        if text_frames_count == 0 and tables_count == 0 and pictures_count == 0:
            issues.append("Slide rỗng: Không tìm thấy bất kỳ đối tượng hiển thị nào.")

        # 5. Cảnh báo thiếu speaker notes
        if not has_speaker_notes:
            issues.append("Thiếu speaker notes (lời dẫn thuyết trình).")

        return {
            "slide_id": slide_id,
            "text_frames_count": text_frames_count,
            "total_text_length": total_text_length,
            "tables_count": tables_count,
            "table_cells_count": table_cells_count,
            "shapes_count": shapes_count,
            "pictures_count": pictures_count,
            "has_speaker_notes": has_speaker_notes,
            "speaker_notes_length": len(speaker_notes_text),
            "is_rasterized": is_rasterized,
            "issues": issues,
        }

    def inspect_deck(self, expected_slide_count: Optional[int] = None) -> Dict[str, Any]:
        """Kiểm tra tổng thể toàn bộ bài thuyết trình đối chiếu với yêu cầu."""
        total_slides = len(self.prs.slides)
        slide_reports: List[Dict[str, Any]] = []
        global_issues: List[str] = []

        # Đối chiếu số lượng slide với yêu cầu
        if expected_slide_count is not None and total_slides != expected_slide_count:
            global_issues.append(
                f"Lệch số lượng slide: Thực tế có {total_slides} slide, nhưng yêu cầu là {expected_slide_count} slide."
            )

        if total_slides == 0:
            global_issues.append("File PPTX không chứa slide nào.")

        has_blocking_errors = False

        for idx, slide in enumerate(self.prs.slides):
            slide_rep = self.inspect_slide(idx, slide)
            slide_reports.append(slide_rep)
            if slide_rep["is_rasterized"]:
                has_blocking_errors = True

        passed = (len(global_issues) == 0) and (not has_blocking_errors)

        return {
            "file_path": str(self.pptx_path),
            "total_slides": total_slides,
            "expected_slides": expected_slide_count,
            "passed": passed,
            "global_issues": global_issues,
            "slides": slide_reports,
        }


def run_inspection(
    pptx_path: str, expected_count: Optional[int] = None, output_json: Optional[str] = None
) -> Dict[str, Any]:
    """Hàm interface dùng để import vào pipeline.py hoặc chạy standalone."""
    inspector = PPTXInspector(pptx_path)
    report = inspector.inspect_deck(expected_slide_count=expected_count)

    if output_json:
        out_p = Path(output_json)
        out_p.parent.mkdir(parents=True, exist_ok=True)
        with open(out_p, "w", encoding="utf-8") as f:
            json.dump(report, f, indent=2, ensure_ascii=False)

    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Kiểm tra cấu trúc, số slide và tính editability của file PPTX."
    )
    parser.add_argument("pptx_file", help="Đường dẫn đến file .pptx cần kiểm tra")
    parser.add_argument(
        "--expected-count",
        type=int,
        default=None,
        help="Số lượng slide kỳ vọng để đối chiếu",
    )
    parser.add_argument(
        "--output",
        "-o",
        default=None,
        help="Đường dẫn file JSON để xuất báo cáo kiểm tra",
    )

    args = parser.parse_args()

    try:
        rep = run_inspection(args.pptx_file, args.expected_count, args.output)
        print(f"\n===== KẾT QUẢ KIỂM TRA PPTX: {args.pptx_file} =====")
        print(f"Tổng số slide: {rep['total_slides']}")
        print(f"Trạng thái đạt chuẩn: {'ĐẠT (PASSED)' if rep['passed'] else 'KHÔNG ĐẠT (FAILED)'}")

        if rep["global_issues"]:
            print("\n[LỖI TỔNG THỂ]:")
            for issue in rep["global_issues"]:
                print(f"  - {issue}")

        blocking_slides = [s for s in rep["slides"] if s["is_rasterized"]]
        if blocking_slides:
            print("\n[CẢNH BÁO CHẶN - SLIDE BỊ RASTER HÓA THÀNH ẢNH]:")
            for s in blocking_slides:
                print(f"  - Slide {s['slide_id']}: {', '.join(s['issues'])}")

        if not rep["passed"]:
            sys.exit(1)
        sys.exit(0)

    except Exception as err:
        print(f"Lỗi khi kiểm tra: {err}", file=sys.stderr)
        sys.exit(2)