"""tests/test_render_deck.py: Unit tests cho render_deck.py."""

import json
from pathlib import Path
from unittest.mock import MagicMock, patch
import pytest
from pptx import Presentation

from render_deck import DeckRenderer, run_render


class TestRenderDeck:
    def test_missing_pptx_raises_filenotfound(self, tmp_path):
        renderer = DeckRenderer()
        with pytest.raises(FileNotFoundError):
            renderer.render_deck(
                pptx_path=tmp_path / "non_existent.pptx",
                output_dir=tmp_path / "renders",
            )

    def test_graceful_degradation_when_no_renderers_available(self, tmp_path):
        pptx_path = tmp_path / "deck.pptx"
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[0])
        prs.save(str(pptx_path))

        renderer = DeckRenderer()
        # Ép buộc bỏ qua GenOffice và kiểm tra môi trường không có libreoffice
        result = renderer.render_deck(
            pptx_path=pptx_path,
            output_dir=tmp_path / "renders",
            force_fallback=True,
        )

        assert result["success"] is False
        assert result["renderer_used"] is None
        assert result["rendered_count"] == 0
        assert len(result["errors"]) > 0

    @patch("render_deck.DeckRenderer._render_genoffice")
    def test_render_genoffice_success_and_normalization(self, mock_genoffice, tmp_path):
        pptx_path = tmp_path / "deck.pptx"
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[0])
        prs.slides.add_slide(prs.slide_layouts[1])
        prs.save(str(pptx_path))

        out_dir = tmp_path / "renders"
        out_dir.mkdir(parents=True, exist_ok=True)

        # Giả lập genoffice tạo 2 file slide_1.png và slide_2.png
        img1 = out_dir / "slide_1.png"
        img2 = out_dir / "slide_2.png"
        img1.write_bytes(b"\x89PNGfake1")
        img2.write_bytes(b"\x89PNGfake2")
        mock_genoffice.return_value = [img1, img2]

        renderer = DeckRenderer()
        result = renderer.render_deck(pptx_path, out_dir)

        assert result["success"] is True
        assert result["renderer_used"] == "genoffice"
        assert result["rendered_count"] == 2
        # Tên file phải được chuẩn hóa thành slide_01.png và slide_02.png
        assert (out_dir / "slide_01.png").exists()
        assert (out_dir / "slide_02.png").exists()

    @patch("render_deck.DeckRenderer._render_genoffice")
    def test_slide_count_mismatch_fails_render(self, mock_genoffice, tmp_path):
        pptx_path = tmp_path / "deck.pptx"
        prs = Presentation()
        # Có 3 slide trong PPTX
        prs.slides.add_slide(prs.slide_layouts[0])
        prs.slides.add_slide(prs.slide_layouts[1])
        prs.slides.add_slide(prs.slide_layouts[1])
        prs.save(str(pptx_path))

        out_dir = tmp_path / "renders"
        out_dir.mkdir(parents=True, exist_ok=True)

        # Nhưng renderer chỉ sinh được 1 ảnh
        img1 = out_dir / "slide_1.png"
        img1.write_bytes(b"\x89PNGfake1")
        mock_genoffice.return_value = [img1]

        renderer = DeckRenderer()
        result = renderer.render_deck(pptx_path, out_dir)

        assert result["success"] is False
        assert any("Lệch số lượng kết xuất" in err for err in result["errors"])

    @patch("render_deck.DeckRenderer._render_genoffice")
    def test_run_render_manifest_output(self, mock_genoffice, tmp_path):
        pptx_path = tmp_path / "deck.pptx"
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[0])
        prs.save(str(pptx_path))

        out_dir = tmp_path / "renders"
        out_dir.mkdir(parents=True, exist_ok=True)
        img1 = out_dir / "slide_01.png"
        img1.write_bytes(b"\x89PNGfake1")
        mock_genoffice.return_value = [img1]

        manifest_file = tmp_path / "render_manifest.json"
        res = run_render(
            pptx_path=str(pptx_path),
            output_dir=str(out_dir),
            output_manifest=str(manifest_file),
        )

        assert res["success"] is True
        assert manifest_file.exists()
        manifest_data = json.loads(manifest_file.read_text(encoding="utf-8"))
        assert manifest_data["rendered_count"] == 1

