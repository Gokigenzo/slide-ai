"""tests/test_inspect_pptx.py: Unit tests cho inspect_pptx.py."""

import json
from pathlib import Path
import pytest
from pptx import Presentation
from pptx.util import Inches

from inspect_pptx import PPTXInspector, run_inspection


class TestInspectPPTX:
    def test_inspect_valid_deck_with_notes_and_tables(self, tmp_path):
        pptx_path = tmp_path / "valid.pptx"
        prs = Presentation()

        # Slide 1: Tiêu đề có speaker notes
        slide1 = prs.slides.add_slide(prs.slide_layouts[0])
        slide1.shapes.title.text = "Tiêu đề bài học"
        slide1.notes_slide.notes_text_frame.text = "Lời chào mở đầu."

        # Slide 2: Nội dung có bảng
        slide2 = prs.slides.add_slide(prs.slide_layouts[5])  # Blank with title
        slide2.shapes.title.text = "Bảng so sánh"
        slide2.notes_slide.notes_text_frame.text = "Phân tích bảng dữ liệu."
        # Thêm bảng native
        shapes = slide2.shapes
        table_shape = shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
        table = table_shape.table
        table.cell(0, 0).text = "Tiêu chí"
        table.cell(0, 1).text = "Giá trị"
        table.cell(1, 0).text = "Tốc độ"
        table.cell(1, 1).text = "Cao"

        prs.save(str(pptx_path))

        report = run_inspection(str(pptx_path), expected_count=2)
        assert report["passed"] is True
        assert report["total_slides"] == 2
        assert len(report["global_issues"]) == 0
        assert report["slides"][0]["has_speaker_notes"] is True
        assert report["slides"][1]["tables_count"] == 1
        assert report["slides"][1]["table_cells_count"] == 4

    def test_detect_slide_count_mismatch(self, tmp_path):
        pptx_path = tmp_path / "count_mismatch.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Slide 1"
        slide.notes_slide.notes_text_frame.text = "Notes"
        prs.save(str(pptx_path))

        report = run_inspection(str(pptx_path), expected_count=5)
        assert report["passed"] is False
        assert any("Lệch số lượng slide" in issue for issue in report["global_issues"])

    def test_detect_missing_speaker_notes(self, tmp_path):
        pptx_path = tmp_path / "no_notes.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Slide không có notes"
        # Không tạo notes_slide
        prs.save(str(pptx_path))

        report = run_inspection(str(pptx_path), expected_count=1)
        slide_rep = report["slides"][0]
        assert slide_rep["has_speaker_notes"] is False
        assert any("Thiếu speaker notes" in issue for issue in slide_rep["issues"])

    def test_detect_rasterized_slide(self, tmp_path):
        pptx_path = tmp_path / "rasterized.pptx"
        prs = Presentation()

        # Tạo ảnh giả
        dummy_img = tmp_path / "dummy.png"
        from PIL import Image
        img = Image.new("RGB", (100, 100), color="blue")
        img.save(str(dummy_img))

        # Slide chỉ có ảnh, không có text box hay bảng
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
        slide.shapes.add_picture(str(dummy_img), Inches(1), Inches(1), Inches(3), Inches(3))
        slide.notes_slide.notes_text_frame.text = "Notes"

        prs.save(str(pptx_path))

        report = run_inspection(str(pptx_path), expected_count=1)
        assert report["passed"] is False
        assert report["slides"][0]["is_rasterized"] is True
        assert any("nghi ngờ raster hóa" in issue for issue in report["slides"][0]["issues"])

    def test_output_json_file(self, tmp_path):
        pptx_path = tmp_path / "deck.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Demo"
        slide.notes_slide.notes_text_frame.text = "Demo notes"
        prs.save(str(pptx_path))

        out_json = tmp_path / "inspection_result.json"
        report = run_inspection(str(pptx_path), output_json=str(out_json))

        assert out_json.exists()
        loaded = json.loads(out_json.read_text(encoding="utf-8"))
        assert loaded["total_slides"] == 1
        assert loaded["passed"] is True

