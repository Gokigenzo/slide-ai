"""tests/test_e2e_workflow.py: Kiểm thử tích hợp toàn diện quy trình 5 giai đoạn của skill slide-ai."""

import json
from pathlib import Path
from unittest.mock import MagicMock, patch
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from inspect_pptx import run_inspection
from pipeline import (
    SlideAIPipeline,
    record_user_approval,
    validate_and_save_request,
    verify_pre_production_integrity,
)
from schema import DeckSpec, SlideSpec


class TestEndToEndWorkflow:
    def test_full_five_phase_lifecycle(self, tmp_path):
        """Kiểm thử trọn vẹn chu trình 5 pha từ Intake đến Final QA."""
        # -------------------------------------------------------------
        # Chuẩn bị dữ liệu bài học nguồn
        # -------------------------------------------------------------
        source_doc = tmp_path / "bai_giang_oop.md"
        source_doc.write_text(
            """# 1. Khái niệm Lập trình Hướng đối tượng
- Định nghĩa: Phương pháp lập trình dựa trên khái niệm các đối tượng
- 4 đặc tính cơ bản: Đóng gói, Kế thừa, Đa hình, Trừu tượng

# 2. Tính Đóng gói (Encapsulation)
- Che giấu trạng thái nội bộ của đối tượng
- Cung cấp getter và setter để truy cập an toàn
- Ví dụ: Lớp BankAccount với số dư private
""",
            encoding="utf-8",
        )

        # =============================================================
        # GIAI ĐOẠN 1: Tiếp nhận và kiểm tra sơ bộ (Intake)
        # =============================================================
        raw_request = {
            "schema_version": 1,
            "operation": "create",
            "project_dir": str(tmp_path),
            "source_files": [
                {"path": "bai_giang_oop.md", "role": "content", "media_type": "md"}
            ],
            "target_deck": None,
            "audience": "Sinh viên năm 2 ngành CNTT",
            "language": "vi",
            "usage": "presentation_and_self_study",
            "slide_count": 4,
            "duration_minutes": 25,
            "style": "academic_minimal",
            "output_name": "oop_lecture.pptx",
            "research_enabled": True,
            "outline_path": "outline.md",
        }

        # Gọi hàm interface chính
        req = validate_and_save_request(tmp_path, raw_request)
        pipe = SlideAIPipeline(project_dir=tmp_path)
        assert pipe.load_state().phase == "sources_validated"

        # =============================================================
        # GIAI ĐOẠN 2: Trích xuất và lập bảng ánh xạ nguồn (Extraction)
        # =============================================================
        manifest = pipe.step_extract_sources()
        assert pipe.manifest_file.exists()
        assert len(manifest.sections) == 2
        assert manifest.sections[0].section_id == "SEC-01"
        assert manifest.sections[1].section_id == "SEC-02"
        assert pipe.load_state().phase == "outline_draft"

        # =============================================================
        # GIAI ĐOẠN 3: Soạn dàn ý và chờ duyệt trên VS Code (Outline Review)
        # =============================================================
        outline_file = pipe.step_generate_outline_draft()
        assert outline_file.exists()
        assert pipe.load_state().phase == "awaiting_outline_approval"

        # Giả lập người dùng mở outline.md trên VS Code và tinh chỉnh nội dung
        user_edited_outline = outline_file.read_text(encoding="utf-8")
        user_edited_outline += "\n<!-- Người dùng bổ sung lưu ý giảng dạy -->\n"
        outline_file.write_text(user_edited_outline, encoding="utf-8")

        # Người dùng chốt phê duyệt
        rec = record_user_approval(tmp_path, explicit_decision=True)
        assert rec.explicit_user_decision is True
        assert pipe.load_state().phase == "approved"

        # =============================================================
        # GIAI ĐOẠN 4: Sản xuất slide qua Adapter (Production)
        # =============================================================
        # Kiểm tra toàn vẹn trước khi cho phép adapter hoạt động
        is_valid, issues = verify_pre_production_integrity(tmp_path)
        assert is_valid is True
        assert len(issues) == 0

        # Chuẩn bị DeckSpec tuân thủ schema
        deck_spec = {
            "schema_version": 1,
            "title": "Bài giảng Lập trình Hướng đối tượng (OOP)",
            "style": "academic_minimal",
            "total_slides": 4,
            "slides": [
                {
                    "slide_id": 1,
                    "section_ids": ["SEC-01", "SEC-02"],
                    "title": "Lập trình Hướng đối tượng",
                    "slide_type": "title",
                    "main_points": ["Tổng quan và các nguyên lý cốt lõi"],
                    "supporting_points": ["Dành cho sinh viên CNTT"],
                    "example_refs": [],
                    "visual_refs": [],
                    "speaker_notes": "Chào các em sinh viên, hôm nay chúng ta sẽ tìm hiểu về OOP.",
                },
                {
                    "slide_id": 2,
                    "section_ids": ["SEC-01"],
                    "title": "Khái niệm OOP & 4 đặc tính",
                    "slide_type": "content",
                    "main_points": [
                        "Định nghĩa: Lập trình xoay quanh khái niệm đối tượng",
                        "4 trụ cột: Đóng gói, Kế thừa, Đa hình, Trừu tượng",
                    ],
                    "supporting_points": ["Nền tảng của Java, C#, C++"],
                    "example_refs": [],
                    "visual_refs": [],
                    "speaker_notes": "Phân tích 4 trụ cột cơ bản của hướng đối tượng.",
                },
                {
                    "slide_id": 3,
                    "section_ids": ["SEC-02"],
                    "title": "Tính Đóng gói (Encapsulation)",
                    "slide_type": "content",
                    "main_points": [
                        "Che giấu dữ liệu nội bộ qua private access modifier",
                        "Cung cấp giao diện an toàn qua getter và setter",
                    ],
                    "supporting_points": ["Tăng tính toàn vẹn và an toàn dữ liệu"],
                    "example_refs": ["Ví dụ: Lớp BankAccount"],
                    "visual_refs": [],
                    "speaker_notes": "Giải thích tại sao không nên để biến public.",
                },
                {
                    "slide_id": 4,
                    "section_ids": ["SEC-01", "SEC-02"],
                    "title": "Tổng kết và Bài tập",
                    "slide_type": "summary",
                    "main_points": [
                        "Nắm vững 4 đặc tính OOP",
                        "Áp dụng tính đóng gói vào thiết kế class thực tế",
                    ],
                    "supporting_points": [],
                    "example_refs": [],
                    "visual_refs": [],
                    "speaker_notes": "Tóm tắt bài học và giao bài tập về nhà.",
                },
            ],
        }

        # Giả lập GenOffice tạo file PPTX chuẩn có native editable text & notes
        output_pptx = pipe.output_dir / "oop_lecture.pptx"

        def mock_build_action(*args, **kwargs):
            prs = Presentation()
            for s_spec in deck_spec["slides"]:
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                slide.shapes.title.text = s_spec["title"]
                tf = slide.shapes.placeholders[1].text_frame
                tf.text = s_spec["main_points"][0]
                slide.notes_slide.notes_text_frame.text = s_spec["speaker_notes"]
            prs.save(str(output_pptx))
            return {"success": True, "output_path": str(output_pptx), "stdout": "Mock build OK"}

        with patch.object(pipe.genoffice_adapter, "preflight", return_value={"available": True, "binary_path": "mock_bin", "error": None}):
            with patch.object(pipe.genoffice_adapter, "build_deck", side_effect=mock_build_action):
                built_pptx = pipe.step_build_deck(deck_spec_data=deck_spec)
                assert built_pptx.exists()
                assert pipe.load_state().phase == "validating"

        # =============================================================
        # GIAI ĐOẠN 5: Kết xuất và kiểm định thị giác (Render & Visual QA)
        # =============================================================
        # 5.1 Kiểm tra cấu trúc hình học qua inspect_pptx
        inspect_report = run_inspection(str(output_pptx), expected_count=4)
        assert inspect_report["passed"] is True
        assert inspect_report["total_slides"] == 4
        # Đảm bảo không có slide nào bị raster hóa
        assert all(not s["is_rasterized"] for s in inspect_report["slides"])
        # Đảm bảo có speaker notes
        assert all(s["has_speaker_notes"] for s in inspect_report["slides"])

        # 5.2 Kết xuất slide thành ảnh
        renders_dir = pipe.renders_dir
        dummy_images = []
        for i in range(1, 5):
            img_p = renders_dir / f"slide_{i:02d}.png"
            img_p.write_bytes(b"\x89PNGmock_rendered_slide")
            dummy_images.append(img_p)

        # 5.3 Thực hiện Visual QA và Finalize
        slide_qa_results = []
        for i in range(1, 5):
            slide_qa_results.append({
                "slide_id": i,
                "readability_passed": True,
                "native_elements_passed": True,
                "rendered_image_path": str(dummy_images[i - 1]),
                "visual_feedback": f"Slide {i} đạt chuẩn bố cục và độ đọc",
                "detected_errors": [],
            })

        final_qa = pipe.step_finalize_qa(slide_qa_results)
        assert final_qa.passed is True
        assert final_qa.total_slides_checked == 4
        assert len(final_qa.blocking_issues) == 0

        # Kiểm tra trạng thái cuối cùng của dự án
        final_state = pipe.load_state()
        assert final_state.phase == "ready"
        assert final_state.user_approved is True
        assert len(final_state.unresolved_issues) == 0

