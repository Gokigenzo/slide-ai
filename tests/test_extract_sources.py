"""tests/test_extract_sources.py: Unit tests cho module trích xuất tài liệu nguồn."""

import json
from pathlib import Path
import pytest
from docx import Document
from pptx import Presentation

from extract_sources import SourceExtractor, compute_sha256, run_extraction
from schema import SourceFile


class TestExtractSources:
    def test_sha256_computation(self, tmp_path):
        f = tmp_path / "sample.txt"
        f.write_text("Hello Slide AI", encoding="utf-8")
        h = compute_sha256(f)
        assert len(h) == 64
        assert isinstance(h, str)

    def test_parse_markdown_sections_and_bullets(self, tmp_path):
        md_file = tmp_path / "content.md"
        content = """# Chương 1: Tổng quan hệ thống
- Điểm chính thứ nhất
- Điểm chính thứ hai
Chi tiết giải thích thêm cho điểm chính có độ dài dài hơn bình thường.

## Chương 2: Kiến trúc dịch vụ
- Microservices kết nối qua gRPC
- Database sử dụng PostgreSQL
"""
        md_file.write_text(content, encoding="utf-8")

        extractor = SourceExtractor(project_dir=tmp_path)
        sections = extractor._parse_markdown(md_file)

        assert len(sections) == 2
        assert sections[0].section_id == "SEC-01"
        assert sections[0].heading == "Chương 1: Tổng quan hệ thống"
        assert "Điểm chính thứ nhất" in sections[0].key_points
        assert "Dòng" in sections[0].locator

        assert sections[1].section_id == "SEC-02"
        assert sections[1].heading == "Chương 2: Kiến trúc dịch vụ"
        assert len(sections[1].key_points) == 2

    def test_parse_docx(self, tmp_path):
        docx_file = tmp_path / "document.docx"
        doc = Document()
        doc.add_heading("1. Giới thiệu mô hình", level=1)
        doc.add_paragraph("Mô hình Client-Server tiêu chuẩn.")
        doc.add_paragraph("Mô hình Peer-to-Peer phân tán.")
        doc.add_heading("2. Giao thức truyền thông", level=1)
        doc.add_paragraph("Sử dụng HTTP/2 và WebSocket.")
        doc.save(str(docx_file))

        extractor = SourceExtractor(project_dir=tmp_path)
        sections = extractor._parse_docx(docx_file)

        assert len(sections) == 2
        assert sections[0].heading == "1. Giới thiệu mô hình"
        assert "Mô hình Client-Server tiêu chuẩn." in sections[0].key_points
        assert sections[1].heading == "2. Giao thức truyền thông"

    def test_parse_pptx_source(self, tmp_path):
        pptx_file = tmp_path / "slides_source.pptx"
        prs = Presentation()
        # Slide 1
        slide1 = prs.slides.add_slide(prs.slide_layouts[0])
        slide1.shapes.title.text = "Bài giảng Mạng máy tính"
        # Slide 2
        slide2 = prs.slides.add_slide(prs.slide_layouts[1])
        slide2.shapes.title.text = "Mô hình OSI"
        tf = slide2.shapes.placeholders[1].text_frame
        tf.text = "Gồm 7 tầng tiêu chuẩn"
        prs.save(str(pptx_file))

        extractor = SourceExtractor(project_dir=tmp_path)
        sections = extractor._parse_pptx(pptx_file)

        assert len(sections) == 2
        assert sections[0].heading == "Bài giảng Mạng máy tính"
        assert sections[1].heading == "Mô hình OSI"
        assert "Gồm 7 tầng tiêu chuẩn" in sections[1].key_points

    def test_extract_manifest_with_visual_reference(self, tmp_path):
        md_file = tmp_path / "main.md"
        md_file.write_text("# Khái niệm\n- Ý cốt lõi", encoding="utf-8")

        ref_file = tmp_path / "diagram.png"
        ref_file.write_bytes(b"\x89PNG\r\n\x1a\nfake_image_bytes")

        extractor = SourceExtractor(project_dir=tmp_path)
        sources = [
            SourceFile(path="main.md", role="content", media_type="md"),
            SourceFile(path="diagram.png", role="visual_reference", media_type="png"),
        ]

        manifest = extractor.extract_manifest(sources)

        assert len(manifest.sources) == 2
        # visual_reference không sinh section nội dung
        assert len(manifest.sections) == 1
        assert manifest.sections[0].heading == "Khái niệm"
        # Cả 2 file đều phải được băm hash
        assert "main.md" in manifest.file_hashes
        assert "diagram.png" in manifest.file_hashes

    def test_missing_source_file_raises(self, tmp_path):
        extractor = SourceExtractor(project_dir=tmp_path)
        sources = [
            SourceFile(path="non_existent.md", role="content", media_type="md")
        ]
        with pytest.raises(FileNotFoundError):
            extractor.extract_manifest(sources)

    def test_run_extraction_cli_interface(self, tmp_path):
        md_file = tmp_path / "lesson.md"
        md_file.write_text("# Bài 1: Nhập môn\n- Nội dung cơ bản", encoding="utf-8")

        req_json = tmp_path / "request.json"
        req_data = {
            "schema_version": 1,
            "operation": "create",
            "project_dir": str(tmp_path),
            "source_files": [
                {"path": "lesson.md", "role": "content", "media_type": "md"}
            ],
            "language": "vi",
            "usage": "presentation_and_self_study",
            "output_name": "test.pptx",
            "research_enabled": True,
            "outline_path": "outline.md",
        }
        req_json.write_text(json.dumps(req_data), encoding="utf-8")

        out_manifest = tmp_path / "source_manifest.json"
        manifest = run_extraction(req_json, out_manifest)

        assert out_manifest.exists()
        assert len(manifest.sections) == 1
        assert manifest.sections[0].heading == "Bài 1: Nhập môn"

